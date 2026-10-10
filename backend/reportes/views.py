from decimal import Decimal
from datetime import datetime, timezone, timedelta
from io import BytesIO


import pymongo
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from django.conf import settings
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone as dj_timezone

from rest_framework import status

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from openpyxl import Workbook

from compras.models import Gasto
from pedidos.models import Pedido, ItemPedido
from produccion.models import ItemProduccionHuevo, RegistroProduccion, Galpon, MovimientoGallina
from users.permissions import IsAdminRole
from pedidos.models import Pedido
from compras.models import Gasto
from produccion.models import RegistroProduccion, Galpon
from .serializers import MetricasComercialesSerializer


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsAdminRole])
def obtener_metricas_comerciales(request):
    """
    Endpoint para calcular los KPIs comerciales conectando 
    con Pedidos y Gastos reales de la base de datos para la US-08.
    """
    hoy = dj_timezone.now()
    mes_actual = hoy.month
    anio_actual = hoy.year
    dias_transcurridos = max(hoy.day, 1)

    # Ventas del mes actual (Pedidos cobrados: estado_pago=True)
    total_ventas = Pedido.objects.filter(
        estado_pago=True,
        fecha_entrega__year=anio_actual,
        fecha_entrega__month=mes_actual
    ).aggregate(total=Sum('total'))['total'] or Decimal('0.00')

    #  Gastos operativos del mes actual (Para calcular la ganancia neta)
    total_gastos = Gasto.objects.filter(
        fecha__year=anio_actual,
        fecha__month=mes_actual
    ).aggregate(total=Sum('monto'))['total'] or Decimal('0.00')

    ganancia_neta = total_ventas - total_gastos

    # Cálculos de Producción y Porcentaje de Postura del mes actual
    registros_mes = RegistroProduccion.objects.filter(
        fecha__year=anio_actual,
        fecha__month=mes_actual
    )
    total_maples_mes = registros_mes.aggregate(total=Sum('total_maples'))['total'] or 0
    
    # Producción diaria promedio de maples
    produccion_diaria_promedio = round( Decimal(total_maples_mes) / Decimal(dias_transcurridos), 2)   # Total de gallinas activas en la granja
    total_gallinas_activas = Galpon.objects.filter(activo=True).aggregate(
        total=Sum('cantidad_actual_gallinas')
    )['total'] or 0

    porcentaje_postura = Decimal('0.00')
    if total_gallinas_activas > 0 and dias_transcurridos > 0:
        total_huevos_mes = total_maples_mes * 30
        
        porcentaje_postura = round(Decimal(total_huevos_mes) / (Decimal(total_gallinas_activas * dias_transcurridos)) * Decimal('100.00'), 2)
        

# 1. Capturar parámetros de fecha opcionales desde el query string (?fecha_desde=...&fecha_hasta=...)
    str_fecha_desde = request.GET.get('fecha_desde')
    str_fecha_hasta = request.GET.get('fecha_hasta')

    tendencia_produccion_meses = []
    nombres_meses_es = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
    nivel_descripcion = "Vista mensual"
    if str_fecha_desde and str_fecha_hasta:
        try:
            f_desde = datetime.strptime(str_fecha_desde, '%Y-%m-%d').date()
            f_hasta = datetime.strptime(str_fecha_hasta, '%Y-%m-%d').date()
            delta_dias = (f_hasta - f_desde).days

            # NIVEL 1: Por Día (Rango corto <= 31 días)
            if delta_dias <= 31:
                nivel_descripcion = "Vista diaria"
                fecha_cursor = f_desde
                while fecha_cursor <= f_hasta:
                    maples_dia = RegistroProduccion.objects.filter(fecha=fecha_cursor).aggregate(total=Sum('total_maples'))['total'] or 0
                    tendencia_produccion_meses.append({
                        'mes': fecha_cursor.strftime('%d/%m'),
                        'promedio': float(maples_dia)
                    })
                    fecha_cursor += timedelta(days=1)

            # NIVEL 2: Por Semana / Intervalos de 7 días (Rango intermedio entre 32 y 120 días)
            elif delta_dias <= 120:
                nivel_descripcion = "Vista semanal"
                fecha_cursor = f_desde
                semana_num = 1
                while fecha_cursor <= f_hasta:
                    fin_semana = min(fecha_cursor + timedelta(days=6), f_hasta)
                    maples_semana = RegistroProduccion.objects.filter(
                        fecha__gte=fecha_cursor,
                        fecha__lte=fin_semana
                    ).aggregate(total=Sum('total_maples'))['total'] or 0
                    
                    tendencia_produccion_meses.append({
                        'mes': fecha_cursor.strftime("%d/%m"),
                        'promedio': float(maples_semana)
                    })
                    fecha_cursor = fin_semana + timedelta(days=1)
                    semana_num += 1

            # NIVEL 3: Por Mes (Rango largo > 120 días)
            else:
                nivel_descripcion = "Vista mensual"
                fecha_cursor = f_desde.replace(day=1)
                while fecha_cursor <= f_hasta:
                    m_year = fecha_cursor.year
                    m_month = fecha_cursor.month
                    maples_m = RegistroProduccion.objects.filter(
                        fecha__year=m_year,
                        fecha__month=m_month
                    ).aggregate(total=Sum('total_maples'))['total'] or 0
                    
                    nombre_etiqueta = f"{nombres_meses_es[m_month - 1]} {m_year}" if m_year != anio_actual else nombres_meses_es[m_month - 1]
                    tendencia_produccion_meses.append({
                        'mes': nombre_etiqueta,
                        'promedio': float(maples_m)
                    })
                    # Avanzar al siguiente mes
                    if m_month == 12:
                        fecha_cursor = fecha_cursor.replace(year=m_year + 1, month=1)
                    else:
                        fecha_cursor = fecha_cursor.replace(month=m_month + 1)

        except ValueError:
            # Fallback por seguridad si el formato de fecha no es válido
            pass

    # Si no se mandaron fechas o falló el parseo, se mantiene el comportamiento por defecto (año actual mes a mes)
    if not tendencia_produccion_meses:
        for m in range(1, mes_actual + 1):
            maples_mes_m = RegistroProduccion.objects.filter(
                fecha__year=anio_actual,
                fecha__month=m
            ).aggregate(total=Sum('total_maples'))['total'] or 0
            
            tendencia_produccion_meses.append({
                'mes': nombres_meses_es[m - 1],
                'promedio': float(maples_mes_m)
            })


    # Datos estructurados para el serializador
    datos_calculados = {
        'ventas_del_mes': total_ventas,
        'porcentaje_cambio_ventas': Decimal('0.00'),
        'produccion_diaria_promedio': produccion_diaria_promedio,
        'porcentaje_postura_mes': porcentaje_postura,
        'ganancia_neta_mensual': ganancia_neta,
        'evolucion_ventas_meses': [{'mes': 'Mes actual', 'total': float(total_ventas)}],
        'tendencia_produccion_meses': tendencia_produccion_meses,
        'nivel_descripcion': nivel_descripcion,
    }

    serializer = MetricasComercialesSerializer(data=datos_calculados)
    serializer.is_valid(raise_exception=True)
    
    return Response(serializer.data, status=status.HTTP_200_OK)


from pathlib import Path
from reportlab.lib.utils import ImageReader

from openpyxl.drawing.image import Image as ExcelImage

def obtener_datos_finanzas(fecha_desde, fecha_hasta):
    ventas = Pedido.objects.filter(estado_pago=True)
    gastos = Gasto.objects.all()

    if fecha_desde:
        ventas = ventas.filter(fecha_entrega__gte=fecha_desde)
        gastos = gastos.filter(fecha__gte=fecha_desde)

    if fecha_hasta:
        ventas = ventas.filter(fecha_entrega__lte=fecha_hasta)
        gastos = gastos.filter(fecha__lte=fecha_hasta)

    total_ingresos = ventas.aggregate(
        total=Sum("total")
    )["total"] or Decimal("0.00")

    total_egresos = gastos.aggregate(
        total=Sum("monto")
    )["total"] or Decimal("0.00")

    balance_neto = total_ingresos - total_egresos

    cantidad_ventas = ventas.count()

    ticket_promedio = (
        (total_ingresos / cantidad_ventas).quantize(Decimal("0.01"))
        if cantidad_ventas > 0
        else Decimal("0.00")
    )

    gastos_por_categoria = (
    gastos
    .values("categoria")
    .annotate(total=Sum("monto"))
    .order_by("categoria")
    )

    ventas_por_tipo = (
    ItemPedido.objects
    .filter(pedido__in=ventas)
    .values("tipo_huevo")
    .annotate(total=Sum("subtotal"))
    .order_by("tipo_huevo")
    )

    return {
        "ingresos": total_ingresos,
        "egresos": total_egresos,
        "balance_neto": balance_neto,
        "cantidad_ventas": cantidad_ventas,
        "ticket_promedio": ticket_promedio,
        "gastos_por_categoria": list(gastos_por_categoria),
        "ventas_por_tipo": list(ventas_por_tipo),
    }

def obtener_datos_produccion(fecha_desde, fecha_hasta):
    registros = RegistroProduccion.objects.all()

    if fecha_desde:
        registros = registros.filter(fecha__gte=fecha_desde)

    if fecha_hasta:
        registros = registros.filter(fecha__lte=fecha_hasta)

    total_maples = registros.aggregate(
        total=Sum("total_maples")
    )["total"] or 0

    total_mermas = registros.aggregate(
        total=Sum("huevos_rotos")
    )["total"] or 0

    items = ItemProduccionHuevo.objects.filter(
        registro__in=registros
    )

    total_huevos = items.aggregate(
    total=Sum("cantidad_huevos")
    )["total"] or 0

    produccion_por_tipo = {}

    for tipo, _ in ItemProduccionHuevo._meta.get_field("tipo_huevo").choices:
        cantidad_huevos = items.filter(
            tipo_huevo=tipo
        ).aggregate(
            total=Sum("cantidad_huevos")
        )["total"] or 0

        produccion_por_tipo[tipo] = (
            cantidad_huevos // ItemProduccionHuevo.HUEVOS_POR_MAPLE
        )

    galpones_activos = Galpon.objects.filter(activo=True)

    total_gallinas = galpones_activos.aggregate(
        total=Sum("cantidad_actual_gallinas")
    )["total"] or 0

    cantidad_galpones_activos = galpones_activos.count()

    movimientos = MovimientoGallina.objects.all()

    if fecha_desde:
        movimientos = movimientos.filter(fecha__gte=fecha_desde)

    if fecha_hasta:
        movimientos = movimientos.filter(fecha__lte=fecha_hasta)

    ingresos_gallinas = movimientos.filter(
        tipo_movimiento=MovimientoGallina.TipoMovimiento.INGRESO
    ).aggregate(
        total=Sum("cantidad_gallinas")
    )["total"] or 0

    salidas_gallinas = movimientos.filter(
        tipo_movimiento=MovimientoGallina.TipoMovimiento.SALIDA
    ).aggregate(
        total=Sum("cantidad_gallinas")
    )["total"] or 0

    mortalidad = movimientos.filter(
        motivo_movimiento=MovimientoGallina.MotivoMovimiento.MUERTE
    ).aggregate(
        total=Sum("cantidad_gallinas")
    )["total"] or 0

    return {
        "total_maples": total_maples,
        "mermas": total_mermas,
        "produccion_por_tipo": produccion_por_tipo,
        "total_gallinas": total_gallinas,
        "galpones_activos": cantidad_galpones_activos,
        "ingresos_gallinas": ingresos_gallinas,
        "salidas_gallinas": salidas_gallinas,
        "mortalidad": mortalidad,
        "total_huevos": total_huevos,
    }

def verificar_espacio_pdf(pdf, y, espacio_necesario=60):
    margen_inferior = 50

    if y - espacio_necesario < margen_inferior:
        pdf.showPage()
        return 800

    return y

def agregar_logo_pdf(pdf):
    ruta_logo = Path(__file__).resolve().parent / "assets" / "logo-granja.png"

    pdf.drawImage(
        ImageReader(str(ruta_logo)),
        480, 755,
        width=60,
        height=60,
        mask="auto",
    )

def agregar_logo_excel(hoja):
    ruta_logo = Path(__file__).resolve().parent / "assets" / "logo-granja.png"

    imagen = ExcelImage(str(ruta_logo))
    imagen.width = 65
    imagen.height = 65

    hoja.add_image(imagen, "D1")

def agregar_finanzas_pdf(pdf, datos, y):
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "RESUMEN FINANCIERO")

    pdf.setFont("Helvetica", 11)
    y -= 25
    pdf.drawString(50, y, f"Ingresos: ${datos['ingresos']}")

    y -= 20
    pdf.drawString(50, y, f"Egresos: ${datos['egresos']}")

    y -= 20
    pdf.drawString(50, y, f"Balance neto: ${datos['balance_neto']}")

    y -= 20
    pdf.drawString(50, y, f"Ventas pagadas: {datos['cantidad_ventas']}")

    y -= 20
    pdf.drawString(50, y, f"Ticket promedio: ${datos['ticket_promedio']}")

    y -= 40
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "INGRESOS POR TIPO DE PRODUCTO")

    pdf.setFont("Helvetica", 11)
    y -= 25

    for venta in datos["ventas_por_tipo"]:
        y = verificar_espacio_pdf(pdf, y)

        nombre_tipo = venta["tipo_huevo"].replace("_", " ").title()
        pdf.drawString(50, y, f"{nombre_tipo}: ${venta['total']}")
        y -= 20

    y -= 20
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "GASTOS POR CATEGORÍA")

    pdf.setFont("Helvetica", 11)
    y -= 25

    for gasto in datos["gastos_por_categoria"]:
        y = verificar_espacio_pdf(pdf, y)

        nombre_categoria = gasto["categoria"].replace("_", " ").title()
        pdf.drawString(50, y, f"{nombre_categoria}: ${gasto['total']}")
        y -= 20

    return y - 20

def agregar_produccion_pdf(pdf, datos, y):
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "RESUMEN DE PRODUCCIÓN")

    pdf.setFont("Helvetica", 11)
    y -= 25
    pdf.drawString(
        50,
        y,
        f"Total producido: {datos['total_maples']} maples",
    )

    y -= 20
    pdf.drawString(
        50,
        y,
        f"Total de huevos: {datos['total_huevos']}",
    )

    y -= 20
    pdf.drawString(
        50,
        y,
        f"Mermas: {datos['mermas']} huevos",
    )

    # Producción por tipo
    y -= 40
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "PRODUCCIÓN POR TIPO")

    pdf.setFont("Helvetica", 11)
    y -= 25

    for tipo, cantidad in datos["produccion_por_tipo"].items():
        y = verificar_espacio_pdf(pdf, y)

        nombre_tipo = tipo.replace("_", " ").title()
        pdf.drawString(50, y, f"{nombre_tipo}: {cantidad} maples")
        y -= 20

    # Plantel actual
    y -= 20
    y = verificar_espacio_pdf(pdf, y, 100)
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "PLANTEL ACTUAL")

    pdf.setFont("Helvetica", 11)
    y -= 25
    pdf.drawString(
        50,
        y,
        f"Gallinas actuales: {datos['total_gallinas']}",
    )

    y -= 20
    pdf.drawString(
        50,
        y,
        f"Galpones activos: {datos['galpones_activos']}",
    )

    # Movimientos de gallinas
    y -= 40
    y = verificar_espacio_pdf(pdf, y, 100)
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "MOVIMIENTOS DE GALLINAS DEL PERÍODO")

    pdf.setFont("Helvetica", 11)
    y -= 25
    pdf.drawString(
        50,
        y,
        f"Ingresos: {datos['ingresos_gallinas']} gallinas",
    )

    y -= 20
    pdf.drawString(
        50,
        y,
        f"Salidas: {datos['salidas_gallinas']} gallinas",
    )

    y -= 20
    pdf.drawString(
        50,
        y,
        f"Mortalidad: {datos['mortalidad']} gallinas",
    )

    return y - 20

def ajustar_ancho_columnas(hoja):
    for columna in hoja.columns:
        ancho_maximo = 0
        letra_columna = columna[0].column_letter

        for celda in columna:
            if celda.value is not None:
                ancho_maximo = max(
                    ancho_maximo,
                    len(str(celda.value))
                )

        hoja.column_dimensions[letra_columna].width = ancho_maximo + 2

def agregar_finanzas_excel(hoja, datos):

    ajustar_ancho_columnas(hoja)

    # Resumen financiero
    hoja.append(["RESUMEN FINANCIERO"])
    hoja.append(["Concepto", "Valor"])
    hoja.append(["Ingresos", datos["ingresos"]])
    hoja.append(["Egresos", datos["egresos"]])
    hoja.append(["Balance neto", datos["balance_neto"]])
    hoja.append(["Ventas pagadas", datos["cantidad_ventas"]])
    hoja.append(["Ticket promedio", datos["ticket_promedio"]])

    hoja.append([])

    # Ingresos por tipo de producto
    hoja.append(["INGRESOS POR TIPO DE PRODUCTO"])
    hoja.append(["Tipo de producto", "Monto"])

    for venta in datos["ventas_por_tipo"]:
        nombre_tipo = venta["tipo_huevo"].replace("_", " ").title()
        hoja.append([nombre_tipo, venta["total"]])

    hoja.append([])

    # Gastos por categoría
    hoja.append(["GASTOS POR CATEGORÍA"])
    hoja.append(["Categoría", "Monto"])

    for gasto in datos["gastos_por_categoria"]:
        nombre_categoria = gasto["categoria"].replace("_", " ").title()
        hoja.append([nombre_categoria, gasto["total"]])

def agregar_produccion_excel(hoja, datos):
    # Resumen de producción
    hoja.append(["RESUMEN DE PRODUCCIÓN"])
    hoja.append(["Concepto", "Valor"])
    hoja.append(["Total producido", datos["total_maples"]])
    hoja.append(["Total de huevos", datos["total_huevos"]])
    hoja.append(["Mermas (huevos)", datos["mermas"]])

    hoja.append([])

    # Producción por tipo
    hoja.append(["PRODUCCIÓN POR TIPO"])
    hoja.append(["Tipo de producto", "Maples"])

    for tipo, cantidad in datos["produccion_por_tipo"].items():
        nombre_tipo = tipo.replace("_", " ").title()
        hoja.append([nombre_tipo, cantidad])

    hoja.append([])

    # Plantel actual
    hoja.append(["PLANTEL ACTUAL"])
    hoja.append(["Concepto", "Cantidad"])
    hoja.append(["Gallinas actuales", datos["total_gallinas"]])
    hoja.append(["Galpones activos", datos["galpones_activos"]])

    hoja.append([])

    # Movimientos del período
    hoja.append(["MOVIMIENTOS DE GALLINAS DEL PERÍODO"])
    hoja.append(["Movimiento", "Cantidad"])
    hoja.append(["Ingresos", datos["ingresos_gallinas"]])
    hoja.append(["Salidas", datos["salidas_gallinas"]])
    hoja.append(["Mortalidad", datos["mortalidad"]])

def convertir_para_mongo(valor):
    if isinstance(valor, Decimal):
        return float(valor)

    if isinstance(valor, dict):
        return {
            clave: convertir_para_mongo(dato)
            for clave, dato in valor.items()
        }

    if isinstance(valor, list):
        return [convertir_para_mongo(dato) for dato in valor]

    return valor

def registrar_exportacion(
    usuario,
    tipo,
    formato,
    fecha_desde,
    fecha_hasta,
    datos,
):
    cliente = pymongo.MongoClient(settings.MONGO_URI)

    try:
        db = cliente["granjanv_reportes"]
        coleccion = db["exportaciones"]

        datos_mongo = convertir_para_mongo(datos)

        registro = {
            "tipo": tipo,
            "formato": formato,
            "fecha_desde": fecha_desde,
            "fecha_hasta": fecha_hasta,
            "datos": datos_mongo,
            "usuario": str(usuario),
            "fecha_generacion": datetime.now(timezone.utc),
        }

        coleccion.insert_one(registro)

    finally:
        cliente.close()


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def reporte_produccion(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_produccion(fecha_desde, fecha_hasta)

    return Response(datos)

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def reporte_finanzas(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_finanzas(fecha_desde, fecha_hasta)
    return Response(datos)


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def reporte_completo(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    finanzas = obtener_datos_finanzas(fecha_desde, fecha_hasta)
    produccion = obtener_datos_produccion(fecha_desde, fecha_hasta)

    datos = {
        "finanzas": finanzas,
        "produccion": produccion,
    }

    return Response(datos)

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_produccion_pdf(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_produccion(fecha_desde, fecha_hasta)

    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle("Reporte de Producción")

    agregar_logo_pdf(pdf)

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 790, "Reporte de Producción")

    pdf.setFont("Helvetica", 11)
    pdf.drawString(
        50,
        765,
        f"Período: {fecha_desde} al {fecha_hasta}",
    )

    agregar_produccion_pdf(pdf, datos, 720)

    pdf.save()

    registrar_exportacion(
        request.user,
        "produccion",
        "pdf",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        'attachment; filename="reporte_produccion.pdf"'
    )

    return response

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_produccion_excel(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_produccion(fecha_desde, fecha_hasta)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Producción"

    agregar_logo_excel(hoja)

    hoja.append(["Reporte de Producción"])
    hoja.append(["Período", f"{fecha_desde} al {fecha_hasta}"])
    hoja.append([])

    agregar_produccion_excel(hoja, datos)

    ajustar_ancho_columnas(hoja)

    buffer = BytesIO()
    workbook.save(buffer)

    registrar_exportacion(
        request.user,
        "produccion",
        "excel",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="reporte_produccion.xlsx"'
    )

    return response

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_finanzas_pdf(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_finanzas(fecha_desde, fecha_hasta)

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle("Reporte de Finanzas")

    agregar_logo_pdf(pdf)

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 790, "Reporte de Finanzas")

    pdf.setFont("Helvetica", 11)
    pdf.drawString(
        50,
        765,
        f"Período: {fecha_desde} al {fecha_hasta}",
    )

    agregar_finanzas_pdf(pdf, datos, 720)

    pdf.save()

    registrar_exportacion(
        request.user,
        "finanzas",
        "pdf",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        'attachment; filename="reporte_finanzas.pdf"'
    )

    return response


@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_finanzas_excel(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    datos = obtener_datos_finanzas(fecha_desde, fecha_hasta)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Finanzas"

    agregar_logo_excel(hoja)

    hoja.append(["Reporte de Finanzas"])
    hoja.append(["Período", f"{fecha_desde} al {fecha_hasta}"])
    hoja.append([])

    agregar_finanzas_excel(hoja, datos)

    ajustar_ancho_columnas(hoja)

    buffer = BytesIO()
    workbook.save(buffer)

    registrar_exportacion(
        request.user,
        "finanzas",
        "excel",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        'attachment; filename="reporte_finanzas.xlsx"'
    )

    return response

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_completo_pdf(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    finanzas = obtener_datos_finanzas(fecha_desde, fecha_hasta)
    produccion = obtener_datos_produccion(fecha_desde, fecha_hasta)

    datos = {
        "finanzas": finanzas,
        "produccion": produccion,
    }

    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle("INFORME COMPLETO")

    agregar_logo_pdf(pdf)

    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 790, "INFORME COMPLETO")

    pdf.setFont("Helvetica", 11)
    pdf.drawString(
        50,
        765,
        f"Período: {fecha_desde} al {fecha_hasta}",
    )

    y = agregar_finanzas_pdf(pdf, finanzas, 730)
    y = verificar_espacio_pdf(pdf, y, 180)

    pdf.line(50, y, 545, y)
    y -= 25
    agregar_produccion_pdf(pdf, produccion, y)

    pdf.save()

    registrar_exportacion(
        request.user,
        "completo",
        "pdf",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )
    response["Content-Disposition"] = (
        'attachment; filename="informe_completo.pdf"'
    )

    return response

@api_view(["GET"])
@permission_classes([IsAuthenticated, IsAdminRole])
def exportar_completo_excel(request):
    fecha_desde = request.query_params.get("fecha_desde")
    fecha_hasta = request.query_params.get("fecha_hasta")

    finanzas = obtener_datos_finanzas(fecha_desde, fecha_hasta)
    produccion = obtener_datos_produccion(fecha_desde, fecha_hasta)

    datos = {
        "finanzas": finanzas,
        "produccion": produccion,
    }

    workbook = Workbook()

    # Hoja Finanzas
    hoja_finanzas = workbook.active
    hoja_finanzas.title = "Finanzas"

    agregar_logo_excel(hoja_finanzas)

    hoja_finanzas.append(["Informe Completo - Finanzas"])
    hoja_finanzas.append(
        ["Período", f"{fecha_desde} al {fecha_hasta}"]
    )
    hoja_finanzas.append([])

    agregar_finanzas_excel(hoja_finanzas, finanzas)

    # Hoja Producción
    hoja_produccion = workbook.create_sheet("Producción")

    agregar_logo_excel(hoja_produccion)

    hoja_produccion.append(["Informe Completo - Producción"])
    hoja_produccion.append(
        ["Período", f"{fecha_desde} al {fecha_hasta}"]
    )
    hoja_produccion.append([])

    agregar_produccion_excel(hoja_produccion, produccion)

    # Ajustar ancho de columnas
    ajustar_ancho_columnas(hoja_finanzas)
    ajustar_ancho_columnas(hoja_produccion)

    buffer = BytesIO()
    workbook.save(buffer)

    registrar_exportacion(
        request.user,
        "completo",
        "excel",
        fecha_desde,
        fecha_hasta,
        datos,
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="informe_completo.xlsx"'
    )

    return response